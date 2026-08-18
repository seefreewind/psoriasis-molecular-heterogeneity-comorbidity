from __future__ import annotations

import shutil
from dataclasses import dataclass
from pathlib import Path

import matplotlib as mpl
import matplotlib.image as mpimg
import matplotlib.pyplot as plt
from docx import Document
from docx.enum.section import WD_SECTION
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.shared import Inches, Pt, RGBColor
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from PIL import Image


ROOT = Path(__file__).resolve().parents[2]
FIG_ROOT = ROOT / "results" / "figures"
OUT_DIR = ROOT / "manuscript" / "figure_package"
COMPOSITE_DIR = OUT_DIR / "review_composites"
DOCX_PATH = OUT_DIR / "CB_main_figures_word_atlas.docx"


mpl.rcParams.update(
    {
        "font.family": "sans-serif",
        "font.sans-serif": ["Arial", "Helvetica", "DejaVu Sans", "sans-serif"],
        "svg.fonttype": "none",
        "pdf.fonttype": 42,
        "font.size": 7,
        "axes.spines.right": False,
        "axes.spines.top": False,
        "axes.linewidth": 0.8,
        "legend.frameon": False,
        "figure.facecolor": "white",
        "savefig.facecolor": "white",
    }
)


@dataclass(frozen=True)
class Panel:
    label: str
    title: str
    path: Path


@dataclass(frozen=True)
class FigureSpec:
    number: str
    title: str
    message: str
    panels: tuple[Panel, ...]
    status: str
    risk: str


FIGURES = [
    FigureSpec(
        "Figure 1",
        "Study design and transition from unstable discrete endotypes to continuous molecular programs",
        "The study begins with a stability boundary: discrete k = 2 classes were not robust, so the paper moves to frozen continuous programs.",
        (
            Panel("a", "Study flow", FIG_ROOT / "Figure1A_study_flow.png"),
            Panel("b", "Patient and tissue structure", FIG_ROOT / "Figure1B_patient_tissue_structure.png"),
            Panel("c", "Latent scatter / fallback view", FIG_ROOT / "Figure1C_fallback_latent_scatter.png"),
            Panel("d", "Cluster stability", FIG_ROOT / "Figure2B_cluster_stability.png"),
        ),
        "Existing panels available; needs final numbering cleanup before submission.",
        "Figure 1 currently mixes Phase 1 and earlier Figure 2B file names; narrative is correct, but source filenames are not manuscript-final.",
    ),
    FigureSpec(
        "Figure 2",
        "Molecular-program prioritization and independent bulk replication",
        "F1, F2, F6 and F7 are retained as frozen programs with skin-primary or systemic-supportive roles.",
        (
            Panel("a", "Axis evidence matrix", FIG_ROOT / "Figure2A_axis_evidence_matrix.png"),
            Panel("b", "Retained gene programs", FIG_ROOT / "Figure2B_retained_axis_gene_programs.png"),
            Panel("c", "Tissue contribution", FIG_ROOT / "Figure2C_tissue_contribution.png"),
            Panel("d", "External replication", FIG_ROOT / "Figure2D_external_replication_support.png"),
            Panel("e", "Primary-axis schematic", FIG_ROOT / "Figure2E_primary_axis_schematic.png"),
        ),
        "Existing panels available; strongest current molecular-program figure.",
        "Panel density is high; final layout may need panel cropping or a two-row asymmetric design.",
    ),
    FigureSpec(
        "Figure 3",
        "Directional single-cell and spatial contextualization",
        "Frozen programs receive donor-level cellular and spatial support, but remain contextual tissue-state programs.",
        (
            Panel("a", "GSE228421 cell-type localization", FIG_ROOT / "phase2b" / "Figure3A_GSE228421_axis_celltype_localization.png"),
            Panel("b", "GSE228421 donor-paired effects", FIG_ROOT / "phase2b" / "Figure3B_GSE228421_donor_paired_effects.png"),
            Panel("c", "Discovery-replication concordance", FIG_ROOT / "Figure3B_discovery_replication_concordance.png"),
            Panel("d", "Sensitivity analysis", FIG_ROOT / "Figure3C_sensitivity_analysis.png"),
        ),
        "Existing panels available; claim strength should remain directional/contextual.",
        "Spatial support is not a definitive patient-level mechanism; legend must state contextual interpretation.",
    ),
    FigureSpec(
        "Figure 4",
        "Axis-specific genetics fails while overall psoriasis susceptibility remains the genetic layer",
        "F1/F2/F6/F7 do not become genetically anchored endotypes; this motivates the pivot to overall psoriasis susceptibility.",
        (
            Panel("a", "Genetic anchoring design", FIG_ROOT / "phase3a" / "Figure6A_genetic_anchoring_design.png"),
            Panel("b", "Axis enrichment", FIG_ROOT / "phase3a" / "Figure6B_axis_enrichment.png"),
            Panel("c", "MHC sensitivity", FIG_ROOT / "phase3a" / "Figure6C_MHC_sensitivity.png"),
            Panel("d", "Matched null distributions", FIG_ROOT / "phase3a" / "Figure6D_matched_null_distributions.png"),
            Panel("e", "Evidence matrix", FIG_ROOT / "phase3a" / "Figure6E_evidence_matrix.png"),
            Panel("f", "Shared versus axis-specific model", FIG_ROOT / "phase3a" / "Figure6F_shared_vs_axis_specific.png"),
        ),
        "Existing Phase 3A panels available, but manuscript numbering should change from old Figure 6 to final Figure 4.",
        "This is a failure-boundary figure. It must not look like an afterthought or supplemental-only negative result.",
    ),
    FigureSpec(
        "Figure 5",
        "Genome-wide and local shared genetic architecture",
        "Overall psoriasis susceptibility, not molecular axes, shows disease-specific shared architecture with CAD, PsA and IBD patterns.",
        (
            Panel("a-c", "LDSC and restricted LAVA architecture", FIG_ROOT / "phase4a4b" / "Figure5_shared_genetic_architecture.png"),
        ),
        "Composite already available in PNG/PDF/SVG/TIFF.",
        "CAD is the cleanest non-neighbour signal; IBD direction must be framed as heterogeneous and QC-sensitive.",
    ),
    FigureSpec(
        "Figure 6",
        "Restricted regulatory prioritization and layered biological model",
        "SMR/HEIDI and coloc prioritize restricted regulatory candidates, but these do not directly overlap frozen molecular programs.",
        (
            Panel("a-d", "Shared-locus eQTL prioritization", FIG_ROOT / "phase4c" / "Figure_Phase4C_shared_locus_eqtl_prioritization.png"),
            Panel("e-g", "Restricted coloc contextualization", FIG_ROOT / "phase4d" / "Figure7_restricted_coloc_contextualization.png"),
        ),
        "Existing Phase 4C/4D composites available; should be merged or split into final Figure 6 plus Extended Data.",
        "This figure is visually large. A final submission version may need a hero panel plus source-data table rather than two dense composites.",
    ),
]


def set_cell_shading(cell, fill: str) -> None:
    tc_pr = cell._tc.get_or_add_tcPr()
    shd = OxmlElement("w:shd")
    shd.set(qn("w:fill"), fill)
    tc_pr.append(shd)


def set_run_font(run, size=10, bold=False, color=None) -> None:
    run.font.name = "Arial"
    run.font.size = Pt(size)
    run.bold = bold
    if color:
        run.font.color.rgb = RGBColor.from_string(color)


def add_heading(doc: Document, text: str, level: int = 1) -> None:
    p = doc.add_heading(level=level)
    run = p.add_run(text)
    set_run_font(run, size=18 if level == 1 else 13, bold=True, color="0F4D92")


def add_labeled_paragraph(doc: Document, label: str, body: str) -> None:
    p = doc.add_paragraph()
    r = p.add_run(f"{label}: ")
    set_run_font(r, size=9, bold=True, color="0F4D92")
    r2 = p.add_run(body)
    set_run_font(r2, size=9)


def make_composite(spec: FigureSpec) -> Path:
    n = len(spec.panels)
    if n == 1:
        src = spec.panels[0].path
        out = COMPOSITE_DIR / f"{spec.number.replace(' ', '_')}_review_composite.png"
        shutil.copyfile(src, out)
        for suffix in [".pdf", ".svg", ".tiff"]:
            sibling = src.with_suffix(suffix)
            if sibling.exists():
                shutil.copyfile(sibling, out.with_suffix(suffix))
        return out

    cols = 2 if n <= 4 else 3
    rows = (n + cols - 1) // cols
    fig = plt.figure(figsize=(7.2, 2.75 * rows), dpi=300)
    gs = fig.add_gridspec(rows, cols, wspace=0.08, hspace=0.20)
    for idx, panel in enumerate(spec.panels):
        ax = fig.add_subplot(gs[idx // cols, idx % cols])
        img = mpimg.imread(panel.path)
        ax.imshow(img)
        ax.set_axis_off()
        ax.text(
            0.005,
            0.995,
            panel.label,
            transform=ax.transAxes,
            ha="left",
            va="top",
            fontsize=9,
            fontweight="bold",
            color="black",
            bbox={"facecolor": "white", "edgecolor": "none", "alpha": 0.85, "pad": 1.5},
        )
        ax.text(
            0.08,
            0.995,
            panel.title,
            transform=ax.transAxes,
            ha="left",
            va="top",
            fontsize=6.5,
            color="#333333",
            bbox={"facecolor": "white", "edgecolor": "none", "alpha": 0.85, "pad": 1.0},
        )
    for idx in range(n, rows * cols):
        ax = fig.add_subplot(gs[idx // cols, idx % cols])
        ax.set_axis_off()
    out_base = COMPOSITE_DIR / f"{spec.number.replace(' ', '_')}_review_composite"
    fig.savefig(out_base.with_suffix(".png"), dpi=300, bbox_inches="tight")
    fig.savefig(out_base.with_suffix(".pdf"), bbox_inches="tight")
    fig.savefig(out_base.with_suffix(".svg"), bbox_inches="tight")
    fig.savefig(out_base.with_suffix(".tiff"), dpi=600, bbox_inches="tight")
    plt.close(fig)
    return out_base.with_suffix(".png")


def create_docx(composites: dict[str, Path]) -> None:
    doc = Document()
    sec = doc.sections[0]
    sec.top_margin = Inches(0.65)
    sec.bottom_margin = Inches(0.65)
    sec.left_margin = Inches(0.65)
    sec.right_margin = Inches(0.65)

    styles = doc.styles
    styles["Normal"].font.name = "Arial"
    styles["Normal"].font.size = Pt(9)

    title = doc.add_paragraph()
    title.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r = title.add_run("Communications Biology Figure Atlas")
    set_run_font(r, size=22, bold=True, color="0F4D92")
    subtitle = doc.add_paragraph()
    subtitle.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r = subtitle.add_run("Psoriasis molecular programs and multisystem shared genetic architecture")
    set_run_font(r, size=11, color="4D4D4D")

    add_heading(doc, "Figure Storyline", 1)
    story = (
        "Figure narrative type: multi-omics integration, disease-mechanism boundary setting, "
        "and shared genetic architecture. The main figures should move from the study contract "
        "and data scope to the instability of discrete endotypes, the reproducibility of continuous "
        "molecular programs, donor-level cellular and spatial contextualization, the negative boundary "
        "for axis-specific genetics, and then the genome-wide, local and regulatory architecture of "
        "overall psoriasis susceptibility. The key visual argument is not that each molecular program "
        "is a genetic endotype. It is that the tissue-state layer and the inherited-liability layer "
        "are connected but non-equivalent."
    )
    p = doc.add_paragraph(story)
    for run in p.runs:
        set_run_font(run, size=9)

    add_heading(doc, "Main Figure Plan", 1)
    table = doc.add_table(rows=1, cols=5)
    table.style = "Table Grid"
    hdr = table.rows[0].cells
    for i, text in enumerate(["Figure", "Purpose", "Expected message", "Current status", "Main risk"]):
        hdr[i].text = text
        set_cell_shading(hdr[i], "DDEBF7")
        for p in hdr[i].paragraphs:
            for run in p.runs:
                set_run_font(run, size=8, bold=True)
    for spec in FIGURES:
        cells = table.add_row().cells
        values = [spec.number, spec.title, spec.message, spec.status, spec.risk]
        for cell, value in zip(cells, values):
            cell.text = value
            for p in cell.paragraphs:
                for run in p.runs:
                    set_run_font(run, size=7)

    for spec in FIGURES:
        doc.add_section(WD_SECTION.NEW_PAGE)
        add_heading(doc, f"{spec.number}. {spec.title}", 1)
        add_labeled_paragraph(doc, "Core message", spec.message)
        add_labeled_paragraph(doc, "Status", spec.status)
        add_labeled_paragraph(doc, "Reviewer risk", spec.risk)
        img_path = composites[spec.number]
        p = doc.add_paragraph()
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        p.add_run().add_picture(str(img_path), width=Inches(7.1))
        cap = doc.add_paragraph()
        cap.alignment = WD_ALIGN_PARAGRAPH.CENTER
        r = cap.add_run(f"Review composite for {spec.number}. Use original SVG/PDF/TIFF sources for final submission.")
        set_run_font(r, size=8, color="4D4D4D")

    doc.add_section(WD_SECTION.NEW_PAGE)
    add_heading(doc, "Supplementary Figure Plan", 1)
    supp = doc.add_table(rows=1, cols=3)
    supp.style = "Table Grid"
    for i, text in enumerate(["Supplementary figure", "Purpose", "Content"]):
        supp.rows[0].cells[i].text = text
        set_cell_shading(supp.rows[0].cells[i], "DDEBF7")
    rows = [
        ("Supplementary Fig. 1", "Molecular discovery QC", "Complete clustering stability, factor stability, missingness and preprocessing summaries."),
        ("Supplementary Fig. 2", "External replication details", "Full GSE244679/GSE121212/GSE61281/GSE147339 projection and sensitivity outputs."),
        ("Supplementary Fig. 3", "Single-cell/spatial sensitivity", "Donor-level bootstrap, dropout robustness, treatment/timepoint sensitivity and full cell-type panels."),
        ("Supplementary Fig. 4", "Phase 3A genetic no-go details", "Full matched-null, MHC sensitivity and gene-set negative-control diagnostics."),
        ("Supplementary Fig. 5", "LDSC/LAVA QC", "Munge QC, h²/intercepts, local h² screening and LD-reference sensitivity."),
        ("Supplementary Fig. 6", "SMR/coloc sensitivity", "Full Tier A/B/C tables, HEIDI filters, coloc PP3/PP4 details and MAF-proxy flags."),
    ]
    for vals in rows:
        cells = supp.add_row().cells
        for cell, val in zip(cells, vals):
            cell.text = val
            for p in cell.paragraphs:
                for run in p.runs:
                    set_run_font(run, size=8)

    add_heading(doc, "Figure-to-Method Mapping", 1)
    mapping = doc.add_table(rows=1, cols=4)
    mapping.style = "Table Grid"
    for i, text in enumerate(["Figure", "Required analysis", "Input", "Output"]):
        mapping.rows[0].cells[i].text = text
        set_cell_shading(mapping.rows[0].cells[i], "DDEBF7")
    map_rows = [
        ("Figure 1", "Cluster stability and continuous-factor discovery", "E-MTAB-14509 matched skin/blood data", "Study design, k = 2 instability, continuous axis decision"),
        ("Figure 2", "Axis prioritization and bulk replication", "Frozen F1/F2/F6/F7 programs and paired skin replication", "Program evidence matrix and tissue contribution"),
        ("Figure 3", "Donor-level single-cell and spatial contextualization", "GSE228421, GSE173706, GSE225475, GSE202011", "Cellular/spatial directional support"),
        ("Figure 4", "Axis-specific genetic anchoring test", "Frozen molecular programs and psoriasis GWAS", "NO-GO for genetically anchored axes"),
        ("Figure 5", "Overall psoriasis × comorbidity LDSC/LAVA", "GCST90472771 plus frozen comorbidity GWAS", "Genome-wide and local shared architecture"),
        ("Figure 6", "SMR/HEIDI, coloc and transcriptomic contextualization", "GTEx v8, Tier 1/2 loci and frozen molecular programs", "Restricted regulatory candidates and no direct axis overlap"),
    ]
    for vals in map_rows:
        cells = mapping.add_row().cells
        for cell, val in zip(cells, vals):
            cell.text = val
            for p in cell.paragraphs:
                for run in p.runs:
                    set_run_font(run, size=8)

    add_heading(doc, "Weakness Check", 1)
    checks = [
        "Figure 1 establishes the study contract, but old source filenames and final manuscript numbering still need cleanup.",
        "Figure 4 is the key failure-boundary figure and should not be demoted to supplementary material.",
        "Figure 6 currently combines two large Phase 4C/4D composites; the final version may need one hero panel plus two evidence panels.",
        "IBD legends should keep the directional-heterogeneity and QC-sensitive wording, not a global negative-liability conclusion.",
        "Images in this Word file are review previews; final submission should use the original PDF/SVG/TIFF files and source data.",
    ]
    for item in checks:
        p = doc.add_paragraph(style=None)
        p.style = styles["Normal"]
        p.paragraph_format.left_indent = Inches(0.18)
        r = p.add_run(f"- {item}")
        set_run_font(r, size=9)

    doc.save(DOCX_PATH)


def main() -> None:
    OUT_DIR.mkdir(parents=True, exist_ok=True)
    COMPOSITE_DIR.mkdir(parents=True, exist_ok=True)
    missing = [str(panel.path.relative_to(ROOT)) for spec in FIGURES for panel in spec.panels if not panel.path.exists()]
    if missing:
        raise FileNotFoundError("Missing figure inputs:\n" + "\n".join(missing))

    composites = {spec.number: make_composite(spec) for spec in FIGURES}
    create_docx(composites)
    print(f"Wrote {DOCX_PATH.relative_to(ROOT)}")
    print(f"Wrote composites to {COMPOSITE_DIR.relative_to(ROOT)}")
    for path in sorted(COMPOSITE_DIR.glob("Figure_*_review_composite.png")):
        with Image.open(path) as img:
            print(f"{path.name}: {img.size[0]}x{img.size[1]}")


if __name__ == "__main__":
    main()
