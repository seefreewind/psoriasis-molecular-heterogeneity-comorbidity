#!/usr/bin/env python3
from __future__ import annotations

import re
from pathlib import Path

from docx import Document
from docx.enum.section import WD_SECTION
from docx.enum.table import WD_ALIGN_VERTICAL, WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt, RGBColor


ROOT = Path(__file__).resolve().parents[2]
MANUSCRIPT = ROOT / "manuscript"
FIGDIR = ROOT / "results" / "figures" / "communications_biology_final"
REPORTS = ROOT / "reports"

SRC_MD = MANUSCRIPT / "Communications_Biology_main_manuscript_v4.md"
OUT_MD = MANUSCRIPT / "Communications_Biology_main_manuscript_v5.md"
OUT_DOCX = MANUSCRIPT / "Communications_Biology_main_manuscript_v5.docx"


FIGURE_LEGENDS = """## Figure Legends

**Figure 1. Study design and transition from unstable discrete endotypes to continuous molecular programs.** Public transcriptomic, single-cell, spatial, GWAS and eQTL resources were organized into a prespecified sequence. The initial k = 2 discrete endotype analysis did not meet the bootstrap stability threshold, leading to continuous multi-view molecular modeling and downstream frozen-program contextualization.

**Figure 2. Molecular-program prioritization and independent bulk replication.** Evidence for F1, F2, F6 and F7 is summarized across program stability, tissue contribution, internal support, independent paired-skin replication and biological interpretability. F1, F2 and F6 were retained as skin-primary tissue programs, whereas F7 was retained as a systemic-supportive candidate.

**Figure 3. Directional single-cell and spatial contextualization of selected molecular programs.** Frozen program scores were examined in donor-level single-cell summaries and spatial transcriptomic contexts. The figure shows directional cellular and spatial support for the selected programs while preserving contextual rather than definitive cell-state language.

**Figure 4. Axis-specific genetics fails while overall psoriasis susceptibility remains the genetic layer.** Frozen F1, F2, F6 and F7 programs were tested for psoriasis genetic anchoring before comorbidity genetics. None met the predefined support threshold, motivating the downstream use of overall psoriasis susceptibility rather than axis-specific genetic exposures.

**Figure 5. Genome-wide and local shared genetic architecture of psoriasis comorbidity.** LDSC genetic correlation and restricted LAVA local genetic correlation show disease-specific sharing between overall psoriasis susceptibility and comorbid outcomes. CAD provides the cleanest non-neighbor systemic signal, PsA acts as a near-neighbor positive control, and Crohn disease/ulcerative colitis show directionally heterogeneous local architectures.

**Figure 6. Restricted regulatory prioritization supports a layered model of psoriasis biology.** Restricted local-rg signals were narrowed through SMR/HEIDI prioritization and coloc. Supported and suggestive regulatory candidates did not show direct gene-membership overlap with the molecular programs, supporting a model in which tissue-state heterogeneity and inherited multisystem liability are connected but non-equivalent biological layers.
"""


FIGURE_INSERT_AFTER = {
    "Continuous molecular programs capture reproducible psoriasis tissue heterogeneity": ["Figure1", "Figure2"],
    "Selected molecular programs show directional cellular and spatial organization": ["Figure3"],
    "Reproducible molecular programs do not define independent inherited genetic axes": ["Figure4"],
    "Local genetic correlation reveals heterogeneous comorbidity architectures": ["Figure5"],
    "Restricted regulatory prioritization separates shared genetic architecture from shared regulatory signals": ["Figure6"],
}


FIGURE_SHORT_CAPTIONS = {
    "Figure1": "Figure 1. Study design and transition from unstable discrete endotypes to continuous molecular programs.",
    "Figure2": "Figure 2. Molecular-program prioritization and independent bulk replication.",
    "Figure3": "Figure 3. Directional single-cell and spatial contextualization.",
    "Figure4": "Figure 4. Axis-specific genetics fails while overall psoriasis susceptibility remains the genetic layer.",
    "Figure5": "Figure 5. Genome-wide and local shared genetic architecture.",
    "Figure6": "Figure 6. Restricted regulatory prioritization and layered biological model.",
}


def set_cell_shading(cell, fill: str) -> None:
    tc_pr = cell._tc.get_or_add_tcPr()
    shd = OxmlElement("w:shd")
    shd.set(qn("w:fill"), fill)
    tc_pr.append(shd)


def set_cell_width(cell, width_dxa: int) -> None:
    tc_pr = cell._tc.get_or_add_tcPr()
    tc_w = tc_pr.first_child_found_in("w:tcW")
    if tc_w is None:
        tc_w = OxmlElement("w:tcW")
        tc_pr.append(tc_w)
    tc_w.set(qn("w:w"), str(width_dxa))
    tc_w.set(qn("w:type"), "dxa")


def style_run(run, size: float | None = None, bold: bool | None = None, color: str | None = None, italic: bool | None = None) -> None:
    run.font.name = "Calibri"
    run._element.rPr.rFonts.set(qn("w:eastAsia"), "Calibri")
    if size is not None:
        run.font.size = Pt(size)
    if bold is not None:
        run.bold = bold
    if italic is not None:
        run.italic = italic
    if color is not None:
        run.font.color.rgb = RGBColor.from_string(color)


def setup_document() -> Document:
    doc = Document()
    sec = doc.sections[0]
    sec.top_margin = Inches(0.8)
    sec.bottom_margin = Inches(0.8)
    sec.left_margin = Inches(0.8)
    sec.right_margin = Inches(0.8)

    styles = doc.styles
    normal = styles["Normal"]
    normal.font.name = "Calibri"
    normal._element.rPr.rFonts.set(qn("w:eastAsia"), "Calibri")
    normal.font.size = Pt(10.5)
    normal.paragraph_format.line_spacing = 1.08
    normal.paragraph_format.space_after = Pt(6)

    for name, size, color, before, after in [
        ("Heading 1", 15, "2E74B5", 14, 7),
        ("Heading 2", 12.5, "2E74B5", 11, 5),
        ("Heading 3", 11.2, "1F4D78", 8, 4),
    ]:
        style = styles[name]
        style.font.name = "Calibri"
        style._element.rPr.rFonts.set(qn("w:eastAsia"), "Calibri")
        style.font.size = Pt(size)
        style.font.color.rgb = RGBColor.from_string(color)
        style.font.bold = True
        style.paragraph_format.space_before = Pt(before)
        style.paragraph_format.space_after = Pt(after)
    return doc


def replace_figure_legends(md: str) -> str:
    pattern = r"## Figure Legends\n.*?(?=\n## Tables\n)"
    return re.sub(pattern, FIGURE_LEGENDS.strip() + "\n", md, flags=re.S)


def iter_blocks(md: str):
    lines = md.splitlines()
    i = 0
    while i < len(lines):
        line = lines[i]
        if not line.strip():
            i += 1
            continue
        if line.startswith("| ") and i + 1 < len(lines) and lines[i + 1].startswith("|"):
            table_lines = [line, lines[i + 1]]
            i += 2
            while i < len(lines) and lines[i].startswith("|"):
                table_lines.append(lines[i])
                i += 1
            yield "table", table_lines
            continue
        if line.startswith("#"):
            yield "heading", line
            i += 1
            continue
        para = [line]
        i += 1
        while i < len(lines) and lines[i].strip() and not lines[i].startswith("#") and not lines[i].startswith("| "):
            para.append(lines[i])
            i += 1
        yield "paragraph", " ".join(x.strip() for x in para)


def add_markdown_paragraph(doc: Document, text: str, style: str | None = None) -> None:
    p = doc.add_paragraph(style=style)
    p.alignment = WD_ALIGN_PARAGRAPH.LEFT
    parts = re.split(r"(\*\*.*?\*\*|\*.*?\*|`.*?`)", text)
    for part in parts:
        if not part:
            continue
        if part.startswith("**") and part.endswith("**"):
            run = p.add_run(part[2:-2])
            style_run(run, bold=True)
        elif part.startswith("*") and part.endswith("*"):
            run = p.add_run(part[1:-1])
            style_run(run, italic=True)
        elif part.startswith("`") and part.endswith("`"):
            run = p.add_run(part[1:-1])
            style_run(run, size=9.5, color="555555")
        else:
            run = p.add_run(part)
            style_run(run)


def parse_md_table(lines: list[str]) -> list[list[str]]:
    rows = []
    for idx, line in enumerate(lines):
        if idx == 1 and set(line.replace("|", "").replace(" ", "").strip()) <= {"-", ":"}:
            continue
        cells = [c.strip().replace("<br>", "\n") for c in line.strip().strip("|").split("|")]
        rows.append(cells)
    return rows


def add_table(doc: Document, rows: list[list[str]]) -> None:
    if not rows:
        return
    n_cols = max(len(r) for r in rows)
    table = doc.add_table(rows=len(rows), cols=n_cols)
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    table.style = "Table Grid"
    widths = [int(9360 / n_cols)] * n_cols
    if n_cols >= 6:
        presets = [860, 700, 950, 1080, 950, 780, 1250, 1900, 890]
        if n_cols <= len(presets):
            widths = presets[:n_cols]
        else:
            widths = [int(9360 / n_cols)] * n_cols
    elif n_cols == 4:
        widths = [1300, 2300, 2500, 3260]
    for r_idx, row in enumerate(rows):
        for c_idx in range(n_cols):
            cell = table.cell(r_idx, c_idx)
            cell.vertical_alignment = WD_ALIGN_VERTICAL.CENTER
            set_cell_width(cell, widths[c_idx])
            txt = row[c_idx] if c_idx < len(row) else ""
            cell.text = ""
            p = cell.paragraphs[0]
            run = p.add_run(txt)
            style_run(run, size=7.2 if n_cols >= 6 else 8.2, bold=(r_idx == 0))
            if r_idx == 0:
                set_cell_shading(cell, "F2F4F7")
    doc.add_paragraph()


def add_figure(doc: Document, fig_name: str) -> None:
    img = FIGDIR / f"{fig_name}.png"
    if not img.exists():
        return
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = p.add_run()
    run.add_picture(str(img), width=Inches(6.7))
    cap = doc.add_paragraph()
    cap.alignment = WD_ALIGN_PARAGRAPH.LEFT
    run = cap.add_run(FIGURE_SHORT_CAPTIONS[fig_name])
    style_run(run, size=8.7, bold=True, color="333333")


def add_title_block(doc: Document, title: str) -> None:
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.LEFT
    r = p.add_run(title)
    style_run(r, size=18, bold=True, color="0B2545")
    p.paragraph_format.space_after = Pt(8)
    meta = doc.add_paragraph()
    r = meta.add_run("Communications Biology main manuscript draft v5 | latest figures embedded")
    style_run(r, size=9, color="555555")
    meta.paragraph_format.space_after = Pt(10)


def build() -> None:
    md = replace_figure_legends(SRC_MD.read_text(encoding="utf-8"))
    OUT_MD.write_text(md.strip() + "\n", encoding="utf-8")

    doc = setup_document()
    current_h3 = None
    first_heading = True
    in_references = False

    for kind, payload in iter_blocks(md):
        if kind == "heading":
            text = payload.lstrip("#").strip()
            level = len(payload) - len(payload.lstrip("#"))
            if first_heading and level == 1:
                add_title_block(doc, text)
                first_heading = False
                continue
            if level <= 2 and current_h3 and current_h3 in FIGURE_INSERT_AFTER:
                for fig in FIGURE_INSERT_AFTER[current_h3]:
                    add_figure(doc, fig)
                current_h3 = None
            if level == 1:
                doc.add_heading(text, level=1)
            elif level == 2:
                doc.add_heading(text, level=1)
                in_references = text == "References"
            else:
                if current_h3 and current_h3 in FIGURE_INSERT_AFTER:
                    for fig in FIGURE_INSERT_AFTER[current_h3]:
                        add_figure(doc, fig)
                current_h3 = text
                doc.add_heading(text, level=2)
            continue
        if kind == "paragraph":
            add_markdown_paragraph(doc, payload, style=None)
            continue
        if kind == "table":
            add_table(doc, parse_md_table(payload))

    if current_h3 and current_h3 in FIGURE_INSERT_AFTER:
        for fig in FIGURE_INSERT_AFTER[current_h3]:
            add_figure(doc, fig)

    doc.core_properties.title = "Molecular heterogeneity and multisystem shared genetic architecture define distinct biological layers of psoriasis"
    doc.core_properties.subject = "Communications Biology main manuscript draft v5"
    doc.save(OUT_DOCX)

    report = f"""# Latest Communications Biology Main Manuscript Build

Status: `DOCX_BUILT_PENDING_RENDER_QA`

Source Markdown: `{SRC_MD}`

Updated Markdown: `{OUT_MD}`

Output DOCX: `{OUT_DOCX}`

Embedded figures:
- Figure 1: `{FIGDIR / 'Figure1.png'}`
- Figure 2: `{FIGDIR / 'Figure2.png'}`
- Figure 3: `{FIGDIR / 'Figure3.png'}`
- Figure 4: `{FIGDIR / 'Figure4.png'}`
- Figure 5: `{FIGDIR / 'Figure5.png'}`
- Figure 6: `{FIGDIR / 'Figure6.png'}`

Figure legends were updated from placeholder wording to final manuscript-facing captions.
"""
    (REPORTS / "CB_LATEST_MAIN_MANUSCRIPT_BUILD_REPORT.md").write_text(report, encoding="utf-8")


if __name__ == "__main__":
    build()
